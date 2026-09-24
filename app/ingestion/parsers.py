"""File-format parsers: raw bytes -> page-aware text.

Pages are kept separate (rather than concatenated) so that every chunk can be
cited with the page it came from.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import PurePath

from app.core.errors import DocumentParseError, UnsupportedFileTypeError

SUPPORTED_EXTENSIONS: dict[str, str] = {
    ".pdf": "pdf",
    ".txt": "txt",
    ".md": "markdown",
    ".markdown": "markdown",
    ".docx": "docx",
}


@dataclass
class PageText:
    text: str
    page: int | None = None  # 1-based; None for formats without pages


@dataclass
class ParsedDocument:
    file_type: str
    pages: list[PageText] = field(default_factory=list)
    num_pages: int | None = None


def detect_file_type(filename: str) -> str:
    ext = PurePath(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{ext or '(none)'}'. Supported: {supported}",
            details={"extension": ext, "supported": sorted(SUPPORTED_EXTENSIONS)},
        )
    return SUPPORTED_EXTENSIONS[ext]


def parse_document(filename: str, content: bytes) -> ParsedDocument:
    """Dispatch on extension. Raises `DocumentParseError` for malformed input."""
    file_type = detect_file_type(filename)
    if not content:
        raise DocumentParseError(f"'{filename}' is empty.")
    parser = _PARSERS[file_type]
    try:
        return parser(content)
    except DocumentParseError:
        raise
    except Exception as exc:  # third-party parsers raise a zoo of exception types
        raise DocumentParseError(f"Could not parse '{filename}' as {file_type}: {exc}") from exc


def _decode_text(content: bytes) -> str:
    if b"\x00" in content[:4096]:
        raise DocumentParseError("File looks binary, not text.")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _parse_text(content: bytes) -> ParsedDocument:
    return ParsedDocument(file_type="txt", pages=[PageText(_decode_text(content))])


def _parse_markdown(content: bytes) -> ParsedDocument:
    return ParsedDocument(file_type="markdown", pages=[PageText(_decode_text(content))])


def _parse_pdf(content: bytes) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    if reader.is_encrypted:
        try:
            reader.decrypt("")  # many "encrypted" PDFs use an empty user password
        except Exception as exc:
            raise DocumentParseError("PDF is password-protected.") from exc
    pages = [PageText(page.extract_text() or "", page=i + 1) for i, page in enumerate(reader.pages)]
    return ParsedDocument(file_type="pdf", pages=pages, num_pages=len(pages))


def _parse_docx(content: bytes) -> ParsedDocument:
    from docx import Document

    doc = Document(io.BytesIO(content))
    text = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return ParsedDocument(file_type="docx", pages=[PageText(text)])


_PARSERS = {
    "pdf": _parse_pdf,
    "txt": _parse_text,
    "markdown": _parse_markdown,
    "docx": _parse_docx,
}
